import csv
import os
import sys

from docx import Document


def extract_info_from_docx(file_path):
    """
    Parses a single .docx file and extracts specific sections based on headers.
    """
    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
        return None

    # Get full file content
    all_paragraphs = [p.text for p in doc.paragraphs]
    full_content = "\n".join(all_paragraphs)

    # Get filename remainder
    filename = os.path.basename(file_path)
    if filename.startswith("Comments - "):
        name_remainder = filename[len("Comments - ") :]
        # Remove extension if desired, otherwise keep it.
        # Keeping extension for accuracy based on prompt "remainder of the name"
    else:
        name_remainder = filename

    # Data placeholders
    data = {
        "file_path": file_path,
        "full_content": full_content,
        "name_remainder": name_remainder,
        "project_description": "",
        "teachers_comment": "",
        "game_link": "",
        "student_code": "",
        "additional_content": "",
    }

    # Indices of paragraphs that have been assigned to a specific category
    # Used to determine what goes into "Additional Content"
    used_indices = set()

    # Helper to find header index
    def find_header_index(header_text):
        for i, text in enumerate(all_paragraphs):
            if text.strip() == header_text:
                return i
        return -1

    # 1. Project Description
    pd_idx = find_header_index("Project Description:")
    if pd_idx != -1:
        used_indices.add(pd_idx)
        # Capture the paragraph immediately following
        if pd_idx + 1 < len(all_paragraphs):
            data["project_description"] = all_paragraphs[pd_idx + 1]
            used_indices.add(pd_idx + 1)

    # 2. Teacher's Comment
    tc_idx = find_header_index("Teacher's Comment:")
    if tc_idx != -1:
        used_indices.add(tc_idx)
        if tc_idx + 1 < len(all_paragraphs):
            data["teachers_comment"] = all_paragraphs[tc_idx + 1]
            used_indices.add(tc_idx + 1)

    # 3. Game Link
    gl_idx = find_header_index("Game Link:")
    if gl_idx != -1:
        used_indices.add(gl_idx)
        if gl_idx + 1 < len(all_paragraphs):
            data["game_link"] = all_paragraphs[gl_idx + 1]
            used_indices.add(gl_idx + 1)

    # 4. Student Code (Captures EVERYTHING after the header)
    sc_idx = find_header_index("Student Code:")
    if sc_idx != -1:
        used_indices.add(sc_idx)
        code_lines = []
        for i in range(sc_idx + 1, len(all_paragraphs)):
            code_lines.append(all_paragraphs[i])
            used_indices.add(i)
        data["student_code"] = "\n".join(code_lines)

    # 5. Additional Content (Anything not marked as used)
    additional_lines = []
    for i, text in enumerate(all_paragraphs):
        if (
            i not in used_indices and text.strip()
        ):  # Check if not used and not just empty whitespace
            additional_lines.append(text)

    data["additional_content"] = "\n".join(additional_lines)

    return data


def main():
    # Set directory to search: explicit argument or current directory
    search_dir = sys.argv[1] if len(sys.argv) > 1 else "."
    output_csv = "scraped_comments.csv"

    print(f"Searching in: {os.path.abspath(search_dir)}")

    csv_headers = [
        "File Path",
        "Full Content",
        "Name Remainder",
        "Project Description",
        "Teacher's Comment",
        "Game Link",
        "Student Code",
        "Additional Content",
    ]

    found_count = 0

    with open(output_csv, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                h.lower().replace(" ", "_").replace("'", "") for h in csv_headers
            ],
        )

        # Map headers to the keys used in the data dictionary
        # This mapping is slightly manual to match the specific keys in extract_info_from_docx
        # but DictWriter needs keys that match the dictionary passed to writerow.
        # Let's align them simply:

        writer = csv.writer(f)
        writer.writerow(csv_headers)

        for root, dirs, files in os.walk(search_dir):
            for file in files:
                if file.endswith(".docx") and file.startswith("Comments - "):
                    file_path = os.path.join(root, file)
                    print(f"Processing: {file}")

                    extracted_data = extract_info_from_docx(file_path)

                    if extracted_data:
                        writer.writerow(
                            [
                                extracted_data["file_path"],
                                extracted_data["full_content"],
                                extracted_data["name_remainder"],
                                extracted_data["project_description"],
                                extracted_data["teachers_comment"],
                                extracted_data["game_link"],
                                extracted_data["student_code"],
                                extracted_data["additional_content"],
                            ]
                        )
                        found_count += 1

    print(f"\nDone! Processed {found_count} files.")
    print(f"Results saved to: {os.path.abspath(output_csv)}")


if __name__ == "__main__":
    main()
